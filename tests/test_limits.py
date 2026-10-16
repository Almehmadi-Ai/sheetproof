"""Size limits are enforced before (file size, ZIP directory) and while (row budget) a file is
parsed, so an oversized or hostile upload is never fully loaded."""

from __future__ import annotations

import zipfile

import pandas as pd
import pytest
from conftest import ALICE, call


@pytest.fixture
def forbid_parsers(monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("the file must be refused before any parser runs")

    monkeypatch.setattr(pd, "ExcelFile", forbidden)
    monkeypatch.setattr(pd, "read_excel", forbidden)
    monkeypatch.setattr(pd, "read_csv", forbidden)


def _zip_bomb(path, size=64 * 2**20):
    """A tiny .xlsx-shaped archive whose sheet XML inflates to `size` bytes."""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        with zf.open("xl/worksheets/sheet1.xml", "w") as member:
            chunk = b"0" * 2**20
            for _ in range(size // len(chunk)):
                member.write(chunk)


def test_zip_bomb_is_refused_before_parsing(tool, owui, tmp_path, forbid_parsers):
    bomb = tmp_path / "bomb.xlsx"
    _zip_bomb(bomb)
    assert bomb.stat().st_size < 200_000
    fid = owui.upload("alice", "bomb.xlsx", bomb.read_bytes())
    res = call(tool, "inspect_spreadsheet", ALICE, file_id=fid)
    assert "zip bomb" in res["error"]


def test_total_uncompressed_size_is_capped(tool, owui, forbid_parsers):
    df = pd.DataFrame({"a": range(20_000), "b": ["text value"] * 20_000})
    fid = owui.upload("alice", "big.xlsx", df=df)
    tool.valves.max_uncompressed_mb = 0.5
    res = call(tool, "inspect_spreadsheet", ALICE, file_id=fid)
    assert "max_uncompressed_mb" in res["error"]


def test_not_a_zip_is_refused(tool, owui, forbid_parsers):
    fid = owui.upload("alice", "fake.xlsx", b"this is not a workbook")
    assert "not a valid Excel workbook" in call(tool, "inspect_spreadsheet", ALICE, file_id=fid)["error"]


@pytest.mark.parametrize("name", ["big.csv", "big.xlsx"])
def test_file_size_is_capped_before_reading(tool, owui, name, monkeypatch):
    fid = owui.upload("alice", name, df=pd.DataFrame({"a": range(1000)}))
    tool.valves.max_file_mb = 0.001
    monkeypatch.setattr(pd, "read_csv", lambda *a, **k: pytest.fail("read_csv ran"))
    monkeypatch.setattr(pd, "ExcelFile", lambda *a, **k: pytest.fail("ExcelFile ran"))
    assert "max_file_mb" in call(tool, "inspect_spreadsheet", ALICE, file_id=fid)["error"]


def test_oversized_csv_is_never_fully_parsed(tool, owui, monkeypatch):
    fid = owui.upload("alice", "big.csv", df=pd.DataFrame({"a": range(5000), "b": range(5000), "c": range(5000)}))
    tool.valves.max_cells = 1000  # 3 columns -> at most 333 rows
    real_read_csv, seen = pd.read_csv, []

    def spy(*args, **kwargs):
        out = real_read_csv(*args, **kwargs)
        seen.append((kwargs.get("nrows"), len(out)))
        return out

    monkeypatch.setattr(pd, "read_csv", spy)
    res = call(tool, "inspect_spreadsheet", ALICE, file_id=fid)
    assert res["error"].startswith("File too large: more than 333 rows x 3 columns")
    assert seen == [(334, 334)]  # parsing stops one row past the budget, not at row 5,000


def test_csv_at_the_limit_loads(tool, owui):
    fid = owui.upload("alice", "ok.csv", df=pd.DataFrame({"a": range(333), "b": range(333), "c": range(333)}))
    tool.valves.max_cells = 1000
    assert call(tool, "inspect_spreadsheet", ALICE, file_id=fid)["rows"] == 333


def test_oversized_xlsx_is_never_fully_parsed(tool, owui, monkeypatch):
    fid = owui.upload("alice", "big.xlsx", df=pd.DataFrame({"a": range(2000), "b": range(2000)}))
    tool.valves.max_cells = 600  # 2 columns -> at most 300 rows
    real_parse, seen = pd.ExcelFile.parse, []

    def spy(self, *args, **kwargs):
        out = real_parse(self, *args, **kwargs)
        seen.append((kwargs.get("nrows"), len(out)))
        return out

    monkeypatch.setattr(pd.ExcelFile, "parse", spy)
    res = call(tool, "inspect_spreadsheet", ALICE, file_id=fid)
    assert res["error"].startswith("File too large: more than 300 rows x 2 columns")
    assert max(n for n, _ in seen) == 301 and max(rows for _, rows in seen) <= 301


def test_declared_sheet_width_bounds_the_row_budget(tool, owui, tmp_path):
    """A sheet whose first rows are narrow but which declares (and uses) many columns further
    down must not get a budget computed from the narrow rows."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 61):
        ws.cell(row=r, column=1, value=r)
        ws.cell(row=r, column=2, value=r * 2)
    ws.cell(row=60, column=200, value="far right")
    path = tmp_path / "wide.xlsx"
    wb.save(path)
    fid = owui.upload("alice", "wide.xlsx", path.read_bytes())
    tool.valves.max_cells = 200 * 50  # 200 declared columns -> 50 rows; the sheet has 59 data rows
    assert "more than 50 rows x 200 columns" in call(tool, "inspect_spreadsheet", ALICE, file_id=fid)["error"]


# --------------------------------------------------------------------------- #
# Every result is capped at max_return_rows (flagged, totals still complete)
# --------------------------------------------------------------------------- #
def test_group_results_are_capped_but_totals_are_complete(tool, sales_id, sales_df):
    tool.valves.max_return_rows = 5
    res = call(
        tool,
        "analyze_spreadsheet",
        ALICE,
        operation="group_aggregate",
        value_column="Units",
        group_by="Order ID",
        file_id=sales_id,
    )
    assert len(res["result"]) == 5
    assert res["summary"]["truncated"] is True and res["summary"]["groups_total"] == 12
    assert res["summary"]["grand_total"] == sales_df["Units"].sum()


@pytest.mark.parametrize(
    ("kwargs", "key"),
    [
        ({"operation": "time_series", "value_column": "Units", "date_column": "Order Date", "freq": "D"}, "result"),
        ({"operation": "compare", "value_column": "Units", "date_column": "Order Date", "freq": "D"}, "result"),
        ({"operation": "compare", "value_column": "Units", "group_by": "Order ID"}, "result"),
        ({"operation": "pivot", "value_column": "Units", "group_by": "Order ID", "columns": ["Region"]}, "result"),
    ],
)
def test_other_results_are_capped(tool, sales_id, kwargs, key):
    tool.valves.max_return_rows = 3
    res = call(tool, "analyze_spreadsheet", ALICE, file_id=sales_id, **kwargs)
    assert res["status"] == "ok", res
    assert len(res[key]) == 3 and res["truncated"] is True


def test_value_counts_are_capped(tool, sales_id):
    tool.valves.max_return_rows = 2
    res = call(
        tool, "analyze_spreadsheet", ALICE, operation="value_counts", value_column="Product", n=100, file_id=sales_id
    )["result"]
    assert len(res["counts"]) == 2 and res["truncated"] is True and res["distinct_values"] == 6


def test_pivot_size_is_checked_before_pivoting(tool, sales_id):
    tool.valves.max_cells = 100  # the file has 84 cells; 12 order ids x 12 dates = 144 pivot cells
    res = call(
        tool,
        "analyze_spreadsheet",
        ALICE,
        operation="pivot",
        value_column="Units",
        group_by="Order ID",
        columns=["Order Date"],
        file_id=sales_id,
    )
    assert res["status"] == "error" and "144 cells" in res["reason"]
    res = call(
        tool,
        "analyze_spreadsheet",
        ALICE,
        operation="transform",
        file_id=sales_id,
        pipeline=[{"op": "pivot", "index": "Order ID", "columns": "Order Date", "values": "Units"}],
    )
    assert res["status"] == "error" and "144 cells" in res["reason"]


def test_pipeline_length_is_bounded(tool, sales_id):
    steps = [{"op": "derive", "name": f"c{i}", "expr": "Units * 2"} for i in range(51)]
    res = call(tool, "analyze_spreadsheet", ALICE, operation="transform", pipeline=steps, file_id=sales_id)
    assert res["status"] == "error" and "at most 50 steps" in res["reason"]
